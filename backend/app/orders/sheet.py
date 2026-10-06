"""Checking an uploaded order sheet (CSV or Excel) row by row before any call starts.

Row numbers are the ones the seller sees in Excel: the header is row 1, the first order is row 2.
"""
import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from app.phone import normalize_mobile

REQUIRED_COLUMNS = ("order_id", "customer_name", "phone", "amount", "product", "address")


class SheetError(Exception):
    """A problem with the whole file (wrong type, unreadable, missing columns)."""


@dataclass
class SheetOrder:
    row: int
    order_id: str
    customer_name: str
    phone: str            # cleaned 10-digit mobile
    amount_paise: int
    product: str
    address: str


@dataclass
class SheetReport:
    orders: list[SheetOrder] = field(default_factory=list)
    problems: list[tuple[int, str]] = field(default_factory=list)  # (row, what's wrong)

    @property
    def ready_count(self) -> int:
        return len(self.orders)

    def summary(self) -> str:
        n = self.ready_count
        return f"{n} order ready" if n == 1 else f"{n} orders ready"

    def problems_text(self) -> str:
        return ", ".join(f"Row {row}: {problem}" for row, problem in self.problems)


def check_sheet(filename: str, data: bytes) -> SheetReport:
    rows = _read_rows(filename, data)
    if not rows:
        raise SheetError("The sheet is empty.")

    header = [_text(c).lower() for c in rows[0]]
    col = {name: header.index(name) for name in REQUIRED_COLUMNS if name in header}
    missing = [name for name in REQUIRED_COLUMNS if name not in col]
    if missing:
        raise SheetError("Missing columns: " + ", ".join(missing) + ".")

    report = SheetReport()
    first_row_of_order: dict[str, int] = {}
    for row_no, row in enumerate(rows[1:], start=2):
        values = {name: _text(row[i]) if i < len(row) else "" for name, i in col.items()}
        if not any(values.values()):
            continue  # blank line

        problems = [f"{name.replace('_', ' ')} missing" for name in REQUIRED_COLUMNS if not values[name]]
        phone = normalize_mobile(values["phone"]) if values["phone"] else None
        if values["phone"] and not phone:
            problems.append("phone invalid")
        amount_paise = _paise(values["amount"]) if values["amount"] else None
        if values["amount"] and amount_paise is None:
            problems.append("amount invalid")
        order_id = values["order_id"]
        if order_id in first_row_of_order:
            problems.append(f"order id repeats row {first_row_of_order[order_id]}")
        elif order_id:
            first_row_of_order[order_id] = row_no

        if problems:
            report.problems.append((row_no, " and ".join(problems)))
        else:
            report.orders.append(SheetOrder(
                row=row_no, order_id=order_id, customer_name=values["customer_name"], phone=phone,
                amount_paise=amount_paise, product=values["product"], address=values["address"],
            ))
    return report


def _read_rows(filename: str, data: bytes) -> list[list]:
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    if name.endswith(".xlsx"):
        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise SheetError("This Excel file can't be opened. Save it again as .xlsx and retry.") from e
        return [list(r) for r in book.active.iter_rows(values_only=True)]
    raise SheetError("Please upload a .csv or .xlsx file.")


def _text(value) -> str:
    """Cell value as clean text; Excel numbers like 9829041122.0 become '9829041122'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _paise(amount: str) -> int | None:
    """'₹1,299', 'Rs. 1299.50' or '1299' → paise; None if not a positive amount."""
    cleaned = amount.replace("₹", "").replace(",", "").replace("INR", "").replace("Rs.", "").replace("Rs", "").strip()
    try:
        rupees = Decimal(cleaned)
    except InvalidOperation:
        return None
    if not rupees.is_finite() or rupees <= 0:
        return None
    return int((rupees * 100).quantize(Decimal(1)))
